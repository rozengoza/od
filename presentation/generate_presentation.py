"""
Generates a PowerPoint summary of the helmet-compliance detection project
(MTech AI, Deep Learning) from whatever training artifacts and screenshots
are available in this folder.

Usage:
    python presentation/generate_presentation.py

Reads from:
    presentation/results/**/results.csv   -- one subfolder per training run
                                              (baseline / cbam), plus whatever
                                              plot images sit alongside it.
                                              Recognized by "baseline"/"cbam"
                                              appearing anywhere in the run
                                              folder's name (case-insensitive) --
                                              works even through Google Drive's
                                              messy nested zip-download folder
                                              names.
    presentation/screenshots/*            -- app demo screenshots (any
                                              image files, any names).
    docs/results_comparison.csv           -- baseline vs CBAM comparison
                                              table, if evaluate.py has
                                              been run.
    data/dataset/*/labels/*.txt           -- class distribution, if the
                                              prepared dataset is present
                                              locally.

Writes to:
    presentation/output/Helmet_Detection_Presentation.pptx

Safe to re-run at any point in the project: sections for artifacts that
don't exist yet (e.g. CBAM results before that run finishes) get a
placeholder slide instead of failing, so you can regenerate a fuller deck
later as more results/screenshots come in.
"""
import csv
from pathlib import Path

import yaml
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

REPO_ROOT = Path(__file__).resolve().parent.parent
PRESENTATION_DIR = REPO_ROOT / "presentation"
RESULTS_DIR = PRESENTATION_DIR / "results"
SCREENSHOTS_DIR = PRESENTATION_DIR / "screenshots"
OUTPUT_DIR = PRESENTATION_DIR / "output"
OUTPUT_PATH = OUTPUT_DIR / "Helmet_Detection_Presentation.pptx"
DATASET_DIR = REPO_ROOT / "data" / "dataset"
COMPARISON_CSV = REPO_ROOT / "docs" / "results_comparison.csv"

TITLE = "Automated Safety-Helmet Compliance Monitoring"
SUBTITLE = "A One-Stage Object Detector (YOLOv8) with a CBAM Attention Head"
PROGRAM = "MTech AI 2024, Kathmandu University"
AUTHOR = "Rozen Shrestha"
COURSE = "Deep Learning"

MUTED = RGBColor(0x60, 0x60, 0x60)
ACCENT = RGBColor(0xC0, 0x39, 0x2B)

CLASS_NAMES = {0: "helmet", 1: "head", 2: "person"}


# --------------------------------------------------------------------------
# Data gathering
# --------------------------------------------------------------------------

def find_runs():
    """Maps 'baseline' / 'cbam' -> run directory, discovered by searching for
    results.csv anywhere under presentation/results/ (handles Drive's nested
    zip-download folder names)."""
    runs = {}
    if not RESULTS_DIR.exists():
        return runs
    for csv_path in RESULTS_DIR.rglob("results.csv"):
        run_dir = csv_path.parent
        lname = run_dir.name.lower()
        if "cbam" in lname:
            key = "cbam"
        elif "baseline" in lname:
            key = "baseline"
        else:
            key = run_dir.name
        if key not in runs or csv_path.stat().st_mtime > (runs[key] / "results.csv").stat().st_mtime:
            runs[key] = run_dir
    return runs


def read_final_metrics(run_dir: Path):
    csv_path = run_dir / "results.csv"
    with open(csv_path, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f, skipinitialspace=True))
    if not rows:
        return None
    last = rows[-1]

    def g(key):
        try:
            return float(last[key])
        except (KeyError, ValueError):
            return None

    return {
        "epoch": last.get("epoch", "?").strip() if isinstance(last.get("epoch"), str) else last.get("epoch"),
        "precision": g("metrics/precision(B)"),
        "recall": g("metrics/recall(B)"),
        "mAP50": g("metrics/mAP50(B)"),
        "mAP50-95": g("metrics/mAP50-95(B)"),
    }


def read_hyperparams(run_dir: Path):
    args_path = run_dir / "args.yaml"
    if not args_path.exists():
        return None
    with open(args_path, encoding="utf-8") as f:
        return yaml.safe_load(f)


def find_image(run_dir: Path, *names):
    for name in names:
        p = run_dir / name
        if p.exists():
            return p
    return None


def read_comparison_csv():
    if not COMPARISON_CSV.exists():
        return None
    with open(COMPARISON_CSV, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def compute_class_counts():
    if not DATASET_DIR.exists():
        return None
    counts = {}
    for split in ("train", "val", "test"):
        split_dir = DATASET_DIR / split / "labels"
        if not split_dir.exists():
            continue
        c = {v: 0 for v in CLASS_NAMES.values()}
        n_images = 0
        for f in split_dir.glob("*.txt"):
            n_images += 1
            for line in f.read_text(encoding="utf-8").splitlines():
                if line.strip():
                    cls_id = int(line.split()[0])
                    c[CLASS_NAMES.get(cls_id, str(cls_id))] += 1
        counts[split] = {"images": n_images, **c}
    return counts


def find_screenshots():
    if not SCREENSHOTS_DIR.exists():
        return []
    exts = {".png", ".jpg", ".jpeg"}
    return sorted(p for p in SCREENSHOTS_DIR.iterdir() if p.is_file() and p.suffix.lower() in exts)


# --------------------------------------------------------------------------
# Slide builders
# --------------------------------------------------------------------------

def add_title_slide(prs, title, subtitle_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    subtitle = slide.placeholders[1].text_frame
    subtitle.clear()
    subtitle.paragraphs[0].text = subtitle_lines[0]
    for line in subtitle_lines[1:]:
        p = subtitle.add_paragraph()
        p.text = line
    return slide


def add_bullet_slide(prs, title, bullets):
    """bullets: list of str (level 0) or (str, level) tuples."""
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, b in enumerate(bullets):
        text, level = b if isinstance(b, tuple) else (b, 0)
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level
    return slide


def _fit_image(image_path, max_width, max_height):
    with PILImage.open(image_path) as im:
        w, h = im.size
    aspect = w / h
    width, height = max_width, Emu(int(max_width / aspect))
    if height > max_height:
        height = max_height
        width = Emu(int(height * aspect))
    return width, height


def add_image_slide(prs, title, image_path, caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    max_width, max_height = Inches(9.0), Inches(5.6)
    width, height = _fit_image(image_path, max_width, max_height)
    left = Emu(int((prs.slide_width - width) / 2))
    top = Inches(1.35)
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    if caption:
        box = slide.shapes.add_textbox(Inches(0.5), Inches(7.05), Inches(9.0), Inches(0.4))
        tf = box.text_frame
        tf.text = caption
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.color.rgb = MUTED
    return slide


def add_two_image_slide(prs, title, left_path, right_path, left_caption=None, right_caption=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    max_w, max_h = Inches(4.5), Inches(5.4)
    top = Inches(1.35)
    for i, (path, caption) in enumerate([(left_path, left_caption), (right_path, right_caption)]):
        if path is None:
            continue
        width, height = _fit_image(path, max_w, max_h)
        left = Inches(0.4) + Inches(4.9) * i
        slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        if caption:
            box = slide.shapes.add_textbox(left, Inches(6.85), max_w, Inches(0.4))
            tf = box.text_frame
            tf.text = caption
            tf.paragraphs[0].font.size = Pt(12)
            tf.paragraphs[0].font.color.rgb = MUTED
    return slide


def add_table_slide(prs, title, headers, rows, note=None):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    n_rows, n_cols = len(rows) + 1, len(headers)
    left, top = Inches(0.5), Inches(1.5)
    width, height = Inches(9.0), Inches(0.5 * n_rows)
    table = slide.shapes.add_table(n_rows, n_cols, left, top, width, height).table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = str(h)
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.text_frame.paragraphs[0].font.size = Pt(13)
    if note:
        box = slide.shapes.add_textbox(left, top + height + Inches(0.2), width, Inches(0.6))
        tf = box.text_frame
        tf.text = note
        tf.paragraphs[0].font.size = Pt(12)
        tf.paragraphs[0].font.italic = True
        tf.paragraphs[0].font.color.rgb = MUTED
    return slide


def add_placeholder_slide(prs, title, message):
    return add_bullet_slide(prs, title, [message])


def fmt(x, digits=3):
    if x is None:
        return "-"
    return f"{x:.{digits}f}"


# --------------------------------------------------------------------------
# Deck assembly
# --------------------------------------------------------------------------

def build_metrics_slides(prs, variant_label, run_dir):
    metrics = read_final_metrics(run_dir)
    hp = read_hyperparams(run_dir) or {}

    rows = [["mAP50", fmt(metrics["mAP50"])], ["mAP50-95", fmt(metrics["mAP50-95"])],
             ["Precision", fmt(metrics["precision"])], ["Recall", fmt(metrics["recall"])]] if metrics else []
    if hp:
        rows.append(["Epochs completed", metrics["epoch"] if metrics else "-"])
        rows.append(["Model / imgsz / batch", f"{hp.get('model', '?')} / {hp.get('imgsz', '?')}px / {hp.get('batch', '?')}"])
    add_table_slide(
        prs, f"Results: {variant_label} — Final Metrics", ["Metric", "Value"], rows,
        note="Overall (mean-across-classes) validation metrics; see the comparison slide for per-class numbers.",
    )

    conf_mat = find_image(run_dir, "confusion_matrix.png")
    results_png = find_image(run_dir, "results.png")
    if conf_mat and results_png:
        add_two_image_slide(
            prs, f"Results: {variant_label} — Training Curves & Confusion Matrix",
            results_png, conf_mat,
            left_caption="Loss / mAP over training epochs", right_caption="Confusion matrix (validation set)",
        )
    elif results_png:
        add_image_slide(prs, f"Results: {variant_label} — Training Curves", results_png)
    elif conf_mat:
        add_image_slide(prs, f"Results: {variant_label} — Confusion Matrix", conf_mat)

    pr_curve = find_image(run_dir, "PR_curve.png", "BoxPR_curve.png")
    pred_sample = find_image(run_dir, "val_batch0_pred.jpg")
    if pr_curve and pred_sample:
        add_two_image_slide(
            prs, f"Results: {variant_label} — Precision-Recall & Sample Predictions",
            pr_curve, pred_sample,
            left_caption="Precision-Recall curve per class", right_caption="Sample validation-batch predictions",
        )
    elif pr_curve:
        add_image_slide(prs, f"Results: {variant_label} — Precision-Recall Curve", pr_curve)
    elif pred_sample:
        add_image_slide(prs, f"Results: {variant_label} — Sample Predictions", pred_sample)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()

    runs = find_runs()
    comparison_rows = read_comparison_csv()
    class_counts = compute_class_counts()
    screenshots = find_screenshots()

    print(f"Found runs: {list(runs.keys()) or 'none'}")
    print(f"Comparison CSV: {'found' if comparison_rows else 'not found'}")
    print(f"Dataset class counts: {'computed' if class_counts else 'not available'}")
    print(f"Screenshots: {len(screenshots)} found")

    # 1. Title
    add_title_slide(prs, TITLE, [SUBTITLE, "", PROGRAM, AUTHOR, COURSE])

    # 2. Agenda
    add_bullet_slide(prs, "Agenda", [
        "Problem statement & real-world motivation",
        "Dataset",
        "Method: one-stage detection + CBAM (novel contribution)",
        "Training setup",
        "Results: baseline vs. CBAM",
        "Real-world implementation (demo app)",
        "Limitations, future work & conclusion",
    ])

    # 3. Problem statement
    add_bullet_slide(prs, "Problem Statement & Motivation", [
        "Helmet non-compliance is a leading, preventable cause of injury on construction sites and among motorcyclists.",
        "Manual monitoring doesn't scale to every gate camera or traffic junction.",
        "Goal: a real-time detector that flags a bare head (no helmet) from a camera feed and logs the violation automatically.",
        ("Target scenario:", 0),
        ("Construction-site entry gate camera, or", 1),
        ("Traffic-junction enforcement camera", 1),
    ])

    # 4. Dataset
    dataset_bullets = [
        "Hard Hat Detection (Kaggle, andrewmvd) — derived from the Safety Helmet Wearing Dataset family",
        "~5,000 real-world images: construction sites, crowds, varying lighting/angle/occlusion",
        "3 classes: helmet, head (= violation), person (context)",
        "Converted from Pascal-VOC XML to YOLO format; split 80% train / 10% val / 10% test",
    ]
    add_bullet_slide(prs, "Dataset", dataset_bullets)

    if class_counts:
        headers = ["Split", "Images", "Helmet boxes", "Head boxes", "Person boxes"]
        rows = [
            [split.capitalize(), c["images"], c["helmet"], c["head"], c["person"]]
            for split, c in class_counts.items()
        ]
        add_table_slide(
            prs, "Dataset — Class Distribution", headers, rows,
            note="Note the helmet/person class imbalance — discussed in Limitations.",
        )
    else:
        add_placeholder_slide(
            prs, "Dataset — Class Distribution",
            "data/dataset/ not found locally — run data/prepare_dataset.py to populate this slide with real counts.",
        )

    baseline_run = runs.get("baseline")
    labels_img = find_image(baseline_run, "labels.jpg") if baseline_run else None
    if labels_img:
        add_image_slide(
            prs, "Dataset — Label Statistics", labels_img,
            caption="Class histogram and box size/position distribution (generated by Ultralytics from the training split).",
        )

    # 5. Method: one-stage detector
    add_bullet_slide(prs, "Method: One-Stage Detection (YOLOv8)", [
        "One-stage: predicts class + box directly per spatial location in a single forward pass",
        ("vs. two-stage (e.g. Faster R-CNN): no separate region-proposal step -> much faster, real-time capable", 1),
        "Backbone -> neck (multi-scale feature fusion, P3/P4/P5) -> 3 detection heads",
        "Losses: CIoU (box), BCE (class), Distribution Focal Loss (precise box edges)",
        "Transfer-learned from COCO-pretrained weights (yolov8s.pt), not trained from scratch",
    ])

    # 6. Method: CBAM (novel contribution)
    add_bullet_slide(prs, "Method: CBAM Attention — the Novel-Method Contribution", [
        "CBAM (Convolutional Block Attention Module, Woo et al. 2018): channel attention + spatial attention",
        "Inserted immediately before each of the 3 detection heads (P3/P4/P5)",
        "Motivation: heads/helmets are small, often-occluded objects in cluttered scenes — attention re-weights features toward them instead of dominant background",
        "Implemented as a standalone, reusable PyTorch module (models/attention.py), spliced into YOLOv8's architecture via a custom YAML (models/yolov8s-cbam.yaml)",
        "Backbone/neck weights still transfer-learned from COCO; only the new CBAM layers start randomly initialized",
    ])

    # 7. Training setup
    hp = read_hyperparams(baseline_run) if baseline_run else None
    if hp:
        rows = [
            ["Model", hp.get("model", "-")],
            ["Image size", f"{hp.get('imgsz', '-')} px"],
            ["Batch size", hp.get("batch", "-")],
            ["Epochs (budget)", hp.get("epochs", "-")],
            ["Optimizer", f"{hp.get('optimizer', '-')} (lr0={hp.get('lr0', '-')}, momentum={hp.get('momentum', '-')})"],
            ["Loss weights", f"box={hp.get('box', '-')}, cls={hp.get('cls', '-')}, dfl={hp.get('dfl', '-')}"],
            ["Early stopping patience", hp.get("patience", "-")],
            ["Hardware", "Google Colab (T4 GPU)"],
        ]
        add_table_slide(prs, "Training Setup", ["Parameter", "Value"], rows)
    else:
        add_bullet_slide(prs, "Training Setup", [
            "60 epochs, 640px, batch 16, YOLOv8s, AdamW (auto-selected)",
            "Transfer-learned from COCO-pretrained weights",
            "Trained on Google Colab (free T4 GPU)",
            "(Run presentation/generate_presentation.py again after copying a results folder with args.yaml for exact logged values.)",
        ])

    # 8-9. Baseline results
    if baseline_run:
        build_metrics_slides(prs, "Baseline (YOLOv8s)", baseline_run)
    else:
        add_placeholder_slide(
            prs, "Results: Baseline",
            "No baseline results found under presentation/results/ — copy the run folder (with results.csv, confusion_matrix.png, etc.) there and re-run this script.",
        )

    # 10-11. CBAM results
    cbam_run = runs.get("cbam")
    if cbam_run:
        build_metrics_slides(prs, "CBAM (YOLOv8s + Attention)", cbam_run)
    else:
        add_placeholder_slide(
            prs, "Results: CBAM (YOLOv8s + Attention)",
            "CBAM training not yet complete / results not yet copied here. Re-run this script once "
            "presentation/results/ contains a run folder with 'cbam' in its name.",
        )

    # 12. Comparison
    if comparison_rows:
        headers = list(comparison_rows[0].keys())
        rows = [[r[h] for h in headers] for r in comparison_rows]
        add_table_slide(
            prs, "Comparison: Baseline vs. CBAM", headers, rows,
            note="Generated by evaluate.py on the held-out test split.",
        )
    else:
        add_placeholder_slide(
            prs, "Comparison: Baseline vs. CBAM",
            "docs/results_comparison.csv not found — run evaluate.py with both checkpoints once CBAM "
            "training finishes, then re-run this script for the full comparison table.",
        )

    # 13. Real-world application
    add_bullet_slide(prs, "Real-World Implementation", [
        "Streamlit application (app/streamlit_app.py) — image, video, or live webcam input",
        "Every detection drawn with class-coded boxes (green = helmet, red = head/violation, blue = person)",
        "Bare-head detections automatically logged: timestamp, confidence, snapshot image, CSV row",
        "Simulates deployment behind a construction-site gate camera or a traffic-enforcement camera",
    ])
    for i, shot in enumerate(screenshots, start=1):
        add_image_slide(prs, f"Application Demo ({i}/{len(screenshots)})", shot)
    if not screenshots:
        add_placeholder_slide(
            prs, "Application Demo",
            "No screenshots found in presentation/screenshots/ — drop image/video-mode screenshots there and re-run this script.",
        )

    # 14. Limitations & future work
    add_bullet_slide(prs, "Limitations & Future Work", [
        "Person class under-represented in the dataset (few labeled instances) -> weak person AP",
        "Trained/evaluated on one public dataset — real deployment footage may differ in angle/lighting",
        "Future: fine-tune on self-recorded footage from the actual deployment camera",
        "Future: ablate CBAM's ratio/kernel_size hyperparameters; try ECA or coordinate attention for comparison",
        "Future: export to ONNX and benchmark on edge hardware for real deployment",
    ])

    # 15. Conclusion
    add_bullet_slide(prs, "Conclusion", [
        "Built a full pipeline: reproducible dataset prep -> transfer-learned YOLOv8 training -> a novel CBAM-attention variant -> quantitative comparison -> a working real-time application",
        "One-stage detector achieves strong helmet/head detection accuracy suitable for real-time monitoring",
        "CBAM attention evaluated as a genuine architectural contribution, not just asserted",
        "Delivered as a working, testable safety-compliance monitoring tool, not just a notebook metric",
    ])

    # 16. Thank you
    add_title_slide(prs, "Thank You", ["Questions?", "", AUTHOR, PROGRAM])

    prs.save(OUTPUT_PATH)
    print(f"\nSaved presentation to: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
